"""Extract text with stable source locations; never execute uploaded content."""
from dataclasses import dataclass
from pathlib import Path
import re
import zipfile


@dataclass
class Section:
    text: str
    location: str


def parse_file(path: Path, suffix: str) -> list[Section]:
    if suffix not in {'.md', '.txt', '.pdf', '.docx'}:
        raise ValueError('不支持此文件格式，请上传 TXT、Markdown、PDF 或 DOCX。')
    try:
        if suffix in {'.md', '.txt'}:
            raw = path.read_bytes()
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = raw.decode('gb18030')
            sections = [Section(p.strip(), f'第 {i + 1} 段') for i, p in enumerate(re.split(r'\n\s*\n', text)) if p.strip()]
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ValueError('请先解密 PDF 再上传。')
            if len(reader.pages) > 1500:
                raise ValueError('PDF 超过 1500 页，请拆分后上传。')
            sections = [Section(page.extract_text() or '', f'第 {i + 1} 页') for i, page in enumerate(reader.pages)]
            unreadable = [s.location for s in sections if not s.text.strip()]
            if unreadable:
                raise ValueError('、'.join(unreadable[:8]) + '未提取到正文，可能为空白或扫描页。为避免遗漏内容，请先 OCR 或移除空白页后重新上传。')
        else:
            from docx import Document
            with zipfile.ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError('DOCX 解压后过大，请拆分后上传。')
            doc = Document(path)
            sections = [Section(p.text, f'第 {i + 1} 段') for i, p in enumerate(doc.paragraphs) if p.text.strip()]
            sections += [Section('\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows), f'表格 {i + 1}') for i, table in enumerate(doc.tables)]
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('文件解析失败，请检查文件是否完整、编码是否正确。') from exc
    sections = [Section(s.text.replace('\x00', '').strip(), s.location) for s in sections if s.text.strip()]
    if not sections:
        raise ValueError('未提取到正文。扫描 PDF 需要先进行 OCR；空文件请补充内容。')
    if sum(len(s.text) for s in sections) > 3_000_000:
        raise ValueError('正文超过 300 万字符，请拆分文件后上传。')
    return sections


def split_sections(sections: list[Section], size: int = 400, overlap: int = 60) -> list[dict]:
    if size < 1 or overlap < 0 or overlap >= size:
        raise ValueError('invalid chunk size or overlap')
    result = []
    for section in sections:
        start = 0
        while start < len(section.text):
            end = min(start + size, len(section.text))
            if end < len(section.text):
                candidates = [section.text.rfind(mark, start + size // 2, end) for mark in ['\n', '。', '！', '？', '. ']]
                boundary = max(candidates)
                if boundary > start:
                    end = boundary + 1
            text = section.text[start:end].strip()
            if text:
                result.append({'text': text, 'location': section.location, 'ordinal': len(result)})
            if end == len(section.text):
                break
            start = max(start + 1, end - overlap)
    return result
