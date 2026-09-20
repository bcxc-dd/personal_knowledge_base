import importlib.util
from pathlib import Path

import pytest


def parser():
    assert importlib.util.find_spec('app.parsing') is not None, 'document parser is missing'
    from app import parsing
    return parsing


def test_long_chinese_text_preserves_end_and_locations():
    p = parser()
    text = '缓存策略需要兼顾一致性。' * 300 + '唯一尾部标记'
    chunks = p.split_sections([p.Section(text, '第 1 段')], size=500, overlap=60)
    assert len(chunks) > 1
    assert all(0 < len(c['text']) <= 500 for c in chunks)
    assert chunks[-1]['text'].endswith('唯一尾部标记')
    assert all(c['location'] == '第 1 段' for c in chunks)


def test_empty_text_is_rejected(tmp_path):
    p = parser()
    path = tmp_path / 'empty.txt'
    path.write_text('  \n ', encoding='utf-8')
    with pytest.raises(ValueError, match='正文'):
        p.parse_file(path, '.txt')


def test_docx_extracts_paragraphs_and_tables(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph('项目上线时间为周五。')
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = '负责人：小林'
    path = tmp_path / 'sample.docx'
    doc.save(path)
    sections = parser().parse_file(path, '.docx')
    assert '周五' in '\n'.join(s.text for s in sections)
    assert '小林' in '\n'.join(s.text for s in sections)


def test_invalid_pdf_and_unsupported_files_report_actionable_errors(tmp_path):
    p = parser()
    path = tmp_path / 'bad.pdf'
    path.write_bytes(b'not a PDF')
    with pytest.raises(ValueError, match='解析'):
        p.parse_file(path, '.pdf')
    with pytest.raises(ValueError, match='格式'):
        p.parse_file(path, '.exe')


def test_mixed_pdf_does_not_silently_skip_unreadable_pages(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 20 250 Td (Visible text) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.add_blank_page(width=300, height=300)
    path = tmp_path / 'mixed.pdf'
    writer.write(path)
    with pytest.raises(ValueError, match='第 2 页'):
        parser().parse_file(path, '.pdf')
